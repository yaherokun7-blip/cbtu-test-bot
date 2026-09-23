import csv
import io
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
ALIASES = {
    "student_id": {"student_id", "รหัสผู้เรียน", "รหัสนักเรียน", "รหัสนักศึกษา"},
    "full_name": {"full_name", "ชื่อ-นามสกุล", "ชื่อ–นามสกุล", "ชื่อ นามสกุล", "ชื่อสกุล"},
    "course": {"course", "หลักสูตร", "หลักสูตร/รุ่น", "รุ่น"},
    "role_name": {"role_name", "ยศ", "ชื่อยศ"},
}


class ImportErrorDetail(ValueError):
    pass


def text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_upload(filename, payload):
    if len(payload) > MAX_BYTES:
        raise ImportErrorDetail("ไฟล์ต้องมีขนาดไม่เกิน 2 MB")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        try:
            content = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ImportErrorDetail("กรุณาบันทึก CSV แบบ UTF-8 หรือใช้ไฟล์ Excel .xlsx")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(content), dialect)
        rows = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise ImportErrorDetail(f"รองรับไม่เกิน {MAX_ROWS:,} แถวต่อครั้ง")
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024 or len(archive.infolist()) > 1000:
                    raise ImportErrorDetail("ไฟล์ Excel ใหญ่หรือซับซ้อนเกินไป กรุณาบันทึกเป็น CSV")
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet.max_column and sheet.max_column > 20:
                    raise ImportErrorDetail("รองรับไม่เกิน 20 คอลัมน์")
                rows = []
                for row in sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=20):
                    if any(cell.data_type == "f" for cell in row):
                        raise ImportErrorDetail("ไฟล์มีสูตร Excel กรุณาคัดลอกแล้ววางเป็นค่าก่อนนำเข้า")
                    values = [cell.value for cell in row]
                    while values and values[-1] is None:
                        values.pop()
                    rows.append(values)
            finally:
                workbook.close()
        except ImportErrorDetail:
            raise
        except Exception as exc:
            raise ImportErrorDetail("อ่านไฟล์ Excel ไม่ได้ กรุณาใช้ไฟล์ .xlsx ที่ไม่ใส่รหัสผ่าน") from exc
    else:
        raise ImportErrorDetail("รองรับเฉพาะ .xlsx และ .csv")

    while rows and not any(text(value) for value in rows[-1]):
        rows.pop()
    if len(rows) < 2:
        raise ImportErrorDetail("ไม่พบข้อมูลผู้เรียนในไฟล์")
    if len(rows) > MAX_ROWS + 1:
        raise ImportErrorDetail(f"รองรับไม่เกิน {MAX_ROWS:,} แถวต่อครั้ง")
    if len(rows[0]) > 20:
        raise ImportErrorDetail("รองรับไม่เกิน 20 คอลัมน์")
    columns = {}
    for index, value in enumerate(rows[0]):
        header = text(value).lower()
        for key, aliases in ALIASES.items():
            if header in aliases:
                if key in columns:
                    raise ImportErrorDetail(f"พบคอลัมน์ซ้ำ: {key}")
                columns[key] = index
    if any(key not in columns for key in ["student_id", "full_name", "course"]):
        raise ImportErrorDetail("ต้องมีคอลัมน์ รหัสผู้เรียน, ชื่อ-นามสกุล และ หลักสูตร (ดาวน์โหลดไฟล์ตัวอย่างได้)")
    results, errors, seen, names = [], [], set(), {}
    for number, row in enumerate(rows[1:], 2):
        if not any(text(value) for value in row):
            continue
        record = {key: text(row[index]) if index < len(row) else "" for key, index in columns.items()}
        record["role_name"] = record.get("role_name") or f"ผู้เรียน • {record['course']}"
        record["row"] = number
        reason = None
        if not all(record[key] for key in ["student_id", "full_name", "course"]):
            reason = "กรอกข้อมูลรหัสผู้เรียน ชื่อ และหลักสูตรให้ครบ"
        elif len(record["student_id"]) > 64 or any(len(record[key]) > 100 for key in ["full_name", "course", "role_name"]):
            reason = "รหัสยาวเกิน 64 ตัวอักษร หรือชื่อ/หลักสูตร/ยศยาวเกิน 100 ตัวอักษร"
        elif any(value.startswith(("=", "+", "@")) for key, value in record.items() if key != "row"):
            reason = "ข้อมูลเริ่มต้นด้วยอักขระสูตร กรุณาตรวจสอบ"
        elif (record["student_id"], record["course"]) in seen:
            reason = "รหัสผู้เรียนและหลักสูตรซ้ำภายในไฟล์"
        elif record["student_id"] in names and names[record["student_id"]] != record["full_name"]:
            reason = "รหัสผู้เรียนเดียวกันมีชื่อไม่ตรงกัน"
        if reason:
            errors.append({"row": number, "message": reason})
        else:
            seen.add((record["student_id"], record["course"]))
            names[record["student_id"]] = record["full_name"]
            results.append(record)
    return results, errors


def csv_bytes(headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        # Spreadsheet applications must not evaluate exported values as formulas.
        writer.writerow([("'" + str(v)) if str(v).lstrip().startswith(("=", "+", "-", "@")) else str(v) for v in row])
    return output.getvalue().encode("utf-8-sig")
